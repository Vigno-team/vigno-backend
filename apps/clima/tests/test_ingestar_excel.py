"""Tests para el comando de ingesta de Excel de estaciones meteorológicas."""

from zoneinfo import ZoneInfo

import openpyxl
import pytest
from django.core.management import call_command
from django.core.management.base import CommandError

from apps.clima.models import MedicionHoraria
from apps.ingesta.models import RegistroCarga

ENCABEZADO_SUPERIOR = [
    "FECHA",
    "TEMPERATURA",
    "TEMPERATURA",
    "TEMPERATURA",
    "HUMEDAD RELATIVA",
    "RAD MAX",
    "VEL. MAX. VIENTO",
    "VEL. MAX. VIENTO",
    "PP",
]
ENCABEZADO_INFERIOR = [
    None,
    "Media diaria",
    "Máxima",
    "Mínima",
    "H.R. Media (%)",
    "(W/m2)",
    "(m/s)",
    "Abejas",
    "(mm)",
]


def fila(fecha, temp=15.0):
    return [fecha, temp, 20.0, 10.0, 50.0, 800.0, 5.0, 100, 0.0]


def fila_vacia(fecha):
    return [fecha] + [None] * 8


def crear_excel(ruta, filas):
    """Excel sintético: 4 filas basura, 2 de encabezado, 1 vacía y luego los datos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "mensual"
    for i in range(4):
        ws.append([f"Basura {i}"])
    ws.append(ENCABEZADO_SUPERIOR)
    ws.append(ENCABEZADO_INFERIOR)
    ws.append([])
    for f in filas:
        ws.append(f)
    wb.save(ruta)
    return str(ruta)


def ingestar(archivo, **kwargs):
    call_command("ingestar_excel", archivo=archivo, estacion="San Clemente", **kwargs)


@pytest.fixture
def excel_sintetico(tmp_path):
    return crear_excel(
        tmp_path / "datos_sinteticos.xlsx",
        [
            # Día del cambio de hora de Chile
            fila("2025-09-07 00:00:00"),
            # Hora 23:00 (error común en los Excel)
            fila("2025-09-08 23:00:00", 16.0),
            # Fecha inválida
            ["00:00", 12.0, 12.0, 12.0, 12.0, 12.0, 12.0, 12, 12.0],
            # Celda de temperatura vacía
            fila("2025-09-10 00:00:00", None),
        ],
    )


@pytest.mark.django_db
def test_ingesta_exitosa_y_reglas_negocio(excel_sintetico):
    ingestar(excel_sintetico)

    carga = RegistroCarga.objects.first()
    assert carga.filas_leidas == 4
    assert carga.filas_rechazadas == 1
    assert carga.filas_aceptadas == 3
    assert "1 columnas de abejas" in carga.columnas_descartadas
    assert "Abejas" in carga.columnas_descartadas
    assert carga.formato_detectado == "B_2023_2026"
    assert carga.estado == "OK"
    assert carga.detalle_rechazos[0]["motivo"] == "Fecha vacía o inválida en una fila con datos"

    # La celda vacía no botó el script y guardó el motivo correcto
    medicion_vacia = MedicionHoraria.objects.get(
        timestamp__date__day=10, variable="temperatura_media"
    )
    assert medicion_vacia.valor is None
    assert "Dato en blanco" in medicion_vacia.motivo_nulo

    # Resolución del cambio de hora
    tz = ZoneInfo("America/Santiago")

    medicion_cambio = MedicionHoraria.objects.filter(timestamp__date__day=7).first()
    assert medicion_cambio.timestamp.astimezone(tz).hour == 12

    medicion_23hrs = MedicionHoraria.objects.filter(timestamp__date__day=8).first()
    assert medicion_23hrs.timestamp.astimezone(tz).hour == 12


@pytest.mark.django_db
def test_ingesta_idempotente(excel_sintetico):
    ingestar(excel_sintetico)
    ingestar(excel_sintetico)

    cargas = RegistroCarga.objects.order_by("id")
    assert cargas.count() == 2

    # La primera vez se inserta todo nuevo
    assert cargas[0].registros_insertados == 21
    assert cargas[0].registros_actualizados == 0

    # La segunda vez no hay nada nuevo ni cambiado
    assert cargas[1].registros_insertados == 0
    assert cargas[1].registros_actualizados == 0
    assert cargas[1].registros_sin_cambio == 21

    assert MedicionHoraria.objects.count() == 21


@pytest.mark.django_db
def test_fecha_repetida_en_el_archivo_deja_la_primera_y_rechaza_las_demas(tmp_path):
    archivo = crear_excel(
        tmp_path / "repetidas.xlsx",
        [
            fila("2025-09-07 00:00:00"),
            fila("2025-09-08 00:00:00", 16.0),
            fila("2025-09-08 23:00:00", 99.0),
        ],
    )

    ingestar(archivo)

    carga = RegistroCarga.objects.get()
    assert carga.filas_leidas == 3
    assert carga.filas_aceptadas == 2
    assert carga.filas_rechazadas == 1
    assert carga.detalle_rechazos == [
        {
            "fila": 10,
            "valor_original": "2025-09-08 23:00:00",
            "motivo": "Fecha repetida en el archivo",
        }
    ]

    assert MedicionHoraria.objects.count() == 14
    dia_8 = MedicionHoraria.objects.get(timestamp__date__day=8, variable="temperatura_media")
    assert dia_8.valor == 16.0  # se quedó la primera


@pytest.mark.django_db
def test_dias_futuros_y_plantilla_vacia_se_rechazan(tmp_path):
    archivo = crear_excel(
        tmp_path / "futuras.xlsx",
        [
            fila("2025-09-07 00:00:00"),
            fila("2025-09-08 00:00:00"),
            fila_vacia("2025-09-09 00:00:00"),
            fila_vacia("2100-01-01 00:00:00"),
        ],
    )

    ingestar(archivo)

    carga = RegistroCarga.objects.get()
    assert carga.filas_leidas == 4
    assert carga.filas_aceptadas == 2
    assert carga.filas_rechazadas == 2
    assert {r["motivo"] for r in carga.detalle_rechazos} == {"Fecha futura o sin datos posteriores"}
    assert MedicionHoraria.objects.count() == 14


@pytest.mark.django_db
def test_temporada_mas_reciente_manda_sin_importar_el_orden(tmp_path):
    f2024 = crear_excel(tmp_path / "t2024.xlsx", [fila("2024-01-15 00:00:00", 30.0)])
    f2023 = crear_excel(tmp_path / "t2023.xlsx", [fila("2024-01-15 00:00:00", 10.0)])
    f2025 = crear_excel(tmp_path / "t2025.xlsx", [fila("2024-01-15 00:00:00", 40.0)])

    ingestar(f2024, temporada=2024)
    ingestar(f2023, temporada=2023)

    medicion = MedicionHoraria.objects.get(variable="temperatura_media")
    assert medicion.valor == 30.0
    assert medicion.temporada == 2024

    carga = RegistroCarga.objects.order_by("id").last()
    assert carga.registros_omitidos == 7
    assert carga.registros_actualizados == 0

    # Una temporada más nueva sí sobrescribe, y solo cuenta lo que cambió
    ingestar(f2025, temporada=2025)
    medicion.refresh_from_db()
    assert medicion.valor == 40.0

    carga = RegistroCarga.objects.order_by("id").last()
    assert carga.registros_actualizados == 1
    assert carga.registros_sin_cambio == 6


@pytest.mark.django_db
def test_valor_no_numerico_se_distingue_del_blanco(tmp_path):
    archivo = crear_excel(tmp_path / "sd.xlsx", [fila("2025-09-07 00:00:00", "s/d")])

    ingestar(archivo)

    medicion = MedicionHoraria.objects.get(variable="temperatura_media")
    assert medicion.valor is None
    assert medicion.motivo_nulo.startswith("Valor no numérico")


@pytest.mark.django_db
def test_falla_si_columna_falta_o_formato_desconocido(tmp_path):
    archivo = tmp_path / "datos_rotos.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "mensual"
    for i in range(4):
        ws.append([f"Basura {i}"])

    # Faltan la mitad de las columnas para forzar el error
    ws.append(["FECHA", "TEMPERATURA"])
    ws.append([None, "Media diaria"])
    ws.append([])
    ws.append(["2025-09-07 00:00:00", 15.0])
    wb.save(archivo)

    with pytest.raises(CommandError, match="no calza con ningun formato"):
        ingestar(str(archivo))


def crear_excel_horario(ruta, horas):
    """Excel con el formato H_prueba de la "Estación horaria de prueba"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "mensual"
    for i in range(4):
        ws.append([f"Basura {i}"])
    ws.append(["FECHA Y HORA", "TEMPERATURA", "HUMEDAD RELATIVA", "PP"])
    ws.append([None, "(°C)", "(%)", "(mm)"])
    ws.append([])
    for hora in horas:
        ws.append([hora, 12.0, 80.0, 0.0])
    wb.save(ruta)
    return str(ruta)


@pytest.mark.django_db
def test_excel_horario_con_cambio_de_hora_rechaza_solo_las_horas_dudosas(tmp_path):
    # En abril las 23:00 del sábado se repiten; en septiembre las 00:00 del domingo no existen.
    archivo = crear_excel_horario(
        tmp_path / "horario.xlsx",
        [
            "2025-04-05 22:00:00",
            "2025-04-05 23:00:00",  # ambigua: viene una sola vez, no se sabe cuál es
            "2025-04-06 00:00:00",
            "2025-09-06 23:00:00",
            "2025-09-07 00:00:00",  # inexistente
            "2025-09-07 01:00:00",
        ],
    )

    call_command(
        "ingestar_excel", archivo=archivo, estacion="Estación horaria de prueba", temporada=2025
    )

    carga = RegistroCarga.objects.get()
    assert carga.estado == "OK"
    assert carga.formato_detectado == "H_prueba"
    assert carga.filas_aceptadas == 4
    assert [(r["fila"], r["motivo"]) for r in carga.detalle_rechazos] == [
        (9, "Hora ambigua por cambio de horario"),
        (12, "Hora ambigua por cambio de horario"),
    ]

    tz = ZoneInfo("America/Santiago")
    guardadas = sorted(
        {
            m.timestamp.astimezone(tz).isoformat()
            for m in MedicionHoraria.objects.filter(frecuencia="H")
        }
    )
    assert guardadas == [
        "2025-04-05T22:00:00-03:00",
        "2025-04-06T00:00:00-04:00",
        "2025-09-06T23:00:00-04:00",
        "2025-09-07T01:00:00-03:00",
    ]


@pytest.mark.django_db
def test_horas_dudosas_aparecen_como_rechazadas_en_el_reporte_de_fallas(tmp_path):
    from rest_framework.test import APIClient

    from apps.clima.models import Estacion

    Estacion.objects.create(nombre="Estación horaria de prueba", codigo="horaria")
    archivo = crear_excel_horario(
        tmp_path / "horario.xlsx",
        ["2025-04-05 22:00:00", "2025-04-05 23:00:00", "2025-04-06 00:00:00"],
    )
    call_command(
        "ingestar_excel", archivo=archivo, estacion="Estación horaria de prueba", temporada=2025
    )

    reporte = (
        APIClient()
        .get("/api/v1/calidad/fallas", {"estacion": "horaria", "temporada": "2024-2025"})
        .json()
    )
    rechazadas = [f for f in reporte["fallas"] if f["tipo"] == "rechazado"]
    assert [(f["fecha"], f["fila"], f["motivo"]) for f in rechazadas] == [
        ("2025-04-05", 9, "Hora ambigua por cambio de horario")
    ]
